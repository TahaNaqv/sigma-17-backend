"""The workbook must carry its own caveat.

An actuary who downloads the reserve workbook and works in Excel never sees the web UI's
disclosure. On the client's production book that mattered: `Selected Method = Reported BF` with
a Reported CDF of 1.00 and a blank Implied LR returns reported claims unchanged, and the sheet
recorded a Bornhuetter-Ferguson that never happened. The `Method Note` column says so in the file.
"""

import tempfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from module1_engine.engine import run_update_reserve_summary

BASE_HEADERS = [
    "Accident_Period", "EP", "Paid Claims", "OS Claims", "Reported Claims", "Reported LR",
]
# Two periods: the first reproduces the production row, the second is healthy.
ROWS = [
    ["2021-Q1", 141_243_571.0, 27_260_563.0, 784_956.0, 28_045_519.0, 0.1985],
    ["2021-Q2", 75_838_585.0, 95_388_480.0, -2_183_000.0, 93_205_480.0, 1.2290],
]


def _workbook(path: Path, *, paid_cdf: float, reported_cdf: float) -> None:
    """A reserve workbook with only what `run_update_reserve_summary` reads."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Reserve Summary"
    ws.append(BASE_HEADERS)
    for row in ROWS:
        ws.append(row)

    for sheet, cdf in (("Paid Claims Triangle", paid_cdf), ("Reported Triangle", reported_cdf)):
        tri = wb.create_sheet(sheet)
        tri.append(["Accident Period", 0, 1])
        tri.append(["Selected LDF", 1, 1])
        # `selected_cdf_row_to_series` reverses the row, so both periods get `cdf`.
        tri.append(["Selected CDF", cdf, cdf])
    wb.save(path)


def _run(tmp: Path, *, paid_cdf=1.0, reported_cdf=1.0, overrides=None) -> dict:
    name = "Fire Payment GROSS 2024-12.xlsx"
    _workbook(tmp / name, paid_cdf=paid_cdf, reported_cdf=reported_cdf)
    run_update_reserve_summary(str(tmp), method_overrides={name: overrides or {}})
    ws = load_workbook(tmp / name, data_only=False)["Reserve Summary"]
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    assert "Method Note" in headers, headers
    note_col = headers.index("Method Note") + 1
    return {
        ws.cell(row=r, column=1).value: ws.cell(row=r, column=note_col).value
        for r in range(2, ws.max_row + 1)
    }


@pytest.fixture
def tmp(tmp_path):
    return tmp_path


def test_the_production_case_is_spelled_out_in_the_sheet(tmp):
    notes = _run(tmp, overrides={
        "2021-Q1": {"selected_method": "Reported BF", "implied_lr": None},
        "2021-Q2": {"selected_method": "Reported BF", "implied_lr": None},
    })
    note = notes["2021-Q1"]
    assert "development term" in note
    assert "expected-loss term" in note
    assert "returns reported claims unchanged" in note
    assert "books no IBNR" in note


def test_a_healthy_row_carries_no_note(tmp):
    notes = _run(tmp, paid_cdf=2.1, reported_cdf=1.6, overrides={
        "2021-Q1": {"selected_method": "Reported BF", "implied_lr": 0.62},
        "2021-Q2": {"selected_method": "Reported BF", "implied_lr": 0.62},
    })
    assert notes["2021-Q1"] is None


def test_the_note_is_appended_last_so_no_formula_moves(tmp):
    name = "Fire Payment GROSS 2024-12.xlsx"
    _workbook(tmp / name, paid_cdf=1.0, reported_cdf=1.0)
    run_update_reserve_summary(str(tmp), method_overrides={name: {}})
    ws = load_workbook(tmp / name)["Reserve Summary"]
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    assert headers[-1] == "Method Note"
    # The formulas address columns by name, so the ones that existed before must be unmoved.
    assert headers[: len(BASE_HEADERS)] == BASE_HEADERS
    assert headers[len(BASE_HEADERS)] == "Implied LR"


def test_the_note_column_is_readable_rather_than_accounting_formatted(tmp):
    name = "Fire Payment GROSS 2024-12.xlsx"
    _workbook(tmp / name, paid_cdf=1.0, reported_cdf=1.0)
    run_update_reserve_summary(str(tmp), method_overrides={
        name: {"2021-Q1": {"selected_method": "Reported BF", "implied_lr": None}}
    })
    ws = load_workbook(tmp / name)["Reserve Summary"]
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    letter = ws.cell(row=1, column=headers.index("Method Note") + 1).column_letter
    assert ws.column_dimensions[letter].width > 30
    assert ws.cell(row=2, column=headers.index("Method Note") + 1).alignment.wrap_text is True
