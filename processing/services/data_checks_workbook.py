"""Render a data-checks report as the client's tables (A.6 summary, A.7 premiums, A.8 claims),
a Basis sheet, and one detail sheet per failing check.

Two artefacts from one renderer:

* **Evidence** (``Data_Checks.xlsx``) — written once at run time into the job's output, so it is
  previewed and retained like every other output. Summary tables plus one sheet per failing
  check listing the discrepant rows (traceable to source file and row).
* **Report** — built on demand from the stored summary and the company's current explanations
  (entered in the app after the run). Summary tables only, so it is instant at any book size.

openpyxl only (a guaranteed dependency; xlsxwriter is optional in production), in write-only
mode: detail sheets can run to a hundred thousand rows each, and cell-by-cell writes made the
reference book alone take eight seconds.
"""

from __future__ import annotations

import io
from typing import Any, Iterable, Mapping

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill

from .data_checks import (
    BASIS_TEXT,
    CATEGORY_CLAIMS,
    CATEGORY_LABELS,
    CATEGORY_PREMIUM,
    DETAIL_ROW_CAP,
    FAIL,
    NOT_RUN,
    PASS,
    DataChecksReport,
)

SUMMARY_SHEET = "Data Checks"
PREMIUM_SHEET = "Premiums Discrepancies"
CLAIMS_SHEET = "Claims Discrepancies"
BASIS_SHEET = "Basis"
EXPLANATION_HEADER = "Company explanation"
#: Row of the summary sheet's column headers (checks start on the next row).
SUMMARY_HEADER_ROW = 4

STATUS_TEXT = {
    PASS: "No discrepancies identified",
    FAIL: "Discrepancies identified",
    NOT_RUN: "Not run",
}
#: Shows amounts in thousands while the cell keeps the exact value: (1,234) for negatives.
THOUSANDS_FORMAT = '#,##0,;(#,##0,);"-"'
COUNT_FORMAT = "#,##0"

_HEADER_FILL = PatternFill("solid", fgColor="4F81BD")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_TITLE_FONT = Font(bold=True, size=13)
_SECTION_FONT = Font(bold=True)
_FAIL_FONT = Font(color="C00000")
_WRAP = Alignment(wrap_text=True, vertical="top")


def detail_sheet_name(check_id: str) -> str:
    return f"Check {check_id}"


def _cell(ws, value, *, font=None, fmt=None, wrap=False, fill=None):
    cell = WriteOnlyCell(ws, value=value)
    if font is not None:
        cell.font = font
    if fmt is not None:
        cell.number_format = fmt
    if wrap:
        cell.alignment = _WRAP
    if fill is not None:
        cell.fill = fill
    return cell


def _header_row(ws, values: Iterable[str]) -> list:
    return [_cell(ws, v, font=_HEADER_FONT, fill=_HEADER_FILL, wrap=True) for v in values]


def _widths(ws, widths: list[int]) -> None:
    from openpyxl.utils import get_column_letter

    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _summary_sheet(wb: Workbook, report: Mapping[str, Any],
                   explanations: Mapping[str, str] | None) -> None:
    ws = wb.create_sheet(SUMMARY_SHEET)
    _widths(ws, [14, 48, 28, 10, 60, 60])
    headers = ["Check category & reference", "Check description", "Discrepancies", "Records", "Basis"]
    if explanations is not None:
        headers.append(EXPLANATION_HEADER)
    ws.freeze_panes = f"A{SUMMARY_HEADER_ROW + 1}"
    ws.append([_cell(ws, "Data Checks Summary", font=_TITLE_FONT)])
    ws.append([f"Valuation date: {report.get('valuation_date') or '-'}"])
    ws.append([])
    ws.append(_header_row(ws, headers))
    category = None
    for check in report["checks"]:
        if check["category"] != category:
            category = check["category"]
            ws.append([_cell(ws, CATEGORY_LABELS[category].upper(), font=_SECTION_FONT)])
        failed = check["status"] == FAIL
        row = [
            check["id"],
            _cell(ws, check["description"], wrap=True),
            _cell(ws, STATUS_TEXT[check["status"]], font=_FAIL_FONT if failed else None),
            None if check["status"] == NOT_RUN else _cell(ws, check["records"], fmt=COUNT_FORMAT),
            _cell(ws, check["basis"], wrap=True),
        ]
        if explanations is not None:
            row.append(_cell(ws, explanations.get(check["id"]) or None, wrap=True))
        ws.append(row)


def _amount_sheet(wb: Workbook, name: str, report: Mapping[str, Any], category: str, title: str,
                  amount_label: str, currency: str) -> None:
    ws = wb.create_sheet(name)
    _widths(ws, [26, 26, 32])
    ws.append([_cell(ws, title, font=_TITLE_FONT)])
    ws.append([])
    ws.append(_header_row(ws, ["Check reference number", "Number of records with discrepancies",
                               f"{amount_label} of discrepancies in {currency} '000"]))
    ws.append([_cell(ws, CATEGORY_LABELS[category].upper(), font=_SECTION_FONT)])
    failed = [c for c in report["checks"] if c["category"] == category and c["status"] == FAIL]
    for check in failed:
        ws.append([check["id"], _cell(ws, check["records"], fmt=COUNT_FORMAT),
                   _cell(ws, check["amount"], fmt=THOUSANDS_FORMAT)])
    if not failed:
        ws.append(["No discrepancies identified"])


def _basis_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet(BASIS_SHEET)
    _widths(ws, [140])
    ws.append([_cell(ws, "Basis of the checks", font=_TITLE_FONT)])
    ws.append([])
    for line in BASIS_TEXT:
        ws.append([_cell(ws, line, wrap=True)])


def _plain(value):
    if value is None or value is pd.NaT:
        return None
    if isinstance(value, float) and pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    return value


def _detail_sheet(wb: Workbook, check_id: str, frame: pd.DataFrame) -> None:
    ws = wb.create_sheet(detail_sheet_name(check_id))
    cols = [str(c) for c in frame.columns]
    _widths(ws, [18] * len(cols))
    ws.freeze_panes = "A3"
    total = len(frame)
    ws.append([f"First {DETAIL_ROW_CAP:,} of {total:,} rows." if total > DETAIL_ROW_CAP
               else f"{total:,} rows."])
    ws.append(_header_row(ws, cols))
    for values in frame.head(DETAIL_ROW_CAP).itertuples(index=False, name=None):
        ws.append([_plain(v) for v in values])


def build_data_checks_workbook(
    report: DataChecksReport | Mapping[str, Any],
    *,
    explanations: Mapping[str, str] | None = None,
    currency: str = "SAR",
) -> bytes:
    """Evidence or report workbook.

    * Live ``DataChecksReport`` -> evidence: summary tables + a detail sheet per failing check.
    * Stored report dict -> summary tables only.
    * ``explanations`` given (even empty) -> the summary carries a Company explanation column.
    """
    live = isinstance(report, DataChecksReport)
    data = report.as_dict() if live else report
    wb = Workbook(write_only=True)
    _summary_sheet(wb, data, explanations)
    _amount_sheet(wb, PREMIUM_SHEET, data, CATEGORY_PREMIUM,
                  "Premiums Data Discrepancies Summary", "Total NWP", currency)
    _amount_sheet(wb, CLAIMS_SHEET, data, CATEGORY_CLAIMS,
                  "Claims Data Discrepancies Summary", "Total net paid claims", currency)
    _basis_sheet(wb)
    if live:
        for result in report.results:
            if result.status == FAIL and result.detail is not None and len(result.detail):
                _detail_sheet(wb, result.id, result.detail)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
