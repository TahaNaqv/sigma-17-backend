"""NaN in engine output must not fail a job that already succeeded.

Production traceback this pins (Cash Flow Allocation, 2026-09-09)::

    psycopg2.errors.InvalidTextRepresentation: invalid input syntax for type json
    DETAIL:  Token "NaN" is invalid.
    CONTEXT:  JSON data, line 1: ...m_ratio": -0.009365161576319028, "exp_ratio": NaN...
      File "/app/processing/tasks.py", line 787, in run_module2_allocate_task
        job.save(update_fields=[...])

The allocate engine had run, produced its workbook and its 47 ULR rows; the job
was then marked FAILED because `Exp Ratio` / `RI %` were blank in the source
Combined_Summary, so those rows carried `float('nan')` — which Postgres refuses
in a json column. The user saw "Module 2 processing failed. Check workbook
formats and required sheets", which pointed at entirely the wrong thing.

These tests run against the real database, because that is the only place the
bug existed: `json.dumps` accepts NaN happily, so a test with tidy numbers or a
mocked save would have passed throughout.
"""

import io
import json
import math

from django.contrib.auth import get_user_model
from django.test import SimpleTestCase, TestCase

from core.jsonsafe import json_safe
from processing.models import Module1Job
from tenants.models import Organization

User = get_user_model()

# One row in the shape the allocate engine emits, with the exact values from the
# production traceback.
PRODUCTION_ROW = {
    "id": "1",
    "reserving_class": "Fire",
    "uwy": "2018",
    "comm_ratio": -0.009365161576319028,
    "exp_ratio": float("nan"),
    "ri_percent": float("nan"),
    "combined_ratio": float("nan"),
    "gwp": 1768576.0,
}


class JsonSafeTests(SimpleTestCase):
    def test_nan_and_infinities_become_none(self):
        self.assertIsNone(json_safe(float("nan")))
        self.assertIsNone(json_safe(float("inf")))
        self.assertIsNone(json_safe(float("-inf")))

    def test_finite_values_are_untouched(self):
        """Including the exact ratio from the traceback — a fix that rounded or
        coerced real numbers would be worse than the bug."""
        for v in (0.0, -0.009365161576319028, 1768576.0, -1.5, 1e300):
            self.assertEqual(json_safe(v), v)

    def test_non_floats_pass_through(self):
        for v in ("Fire", 42, True, None, "2018"):
            self.assertEqual(json_safe(v), v)

    def test_recurses_through_nested_containers(self):
        out = json_safe({"rows": [{"a": float("nan"), "b": [float("inf"), 2.5]}]})
        self.assertEqual(out, {"rows": [{"a": None, "b": [None, 2.5]}]})

    def test_tuples_become_lists(self):
        self.assertEqual(json_safe((1.0, float("nan"))), [1.0, None])

    def test_input_is_not_mutated(self):
        original = {"a": float("nan")}
        json_safe(original)
        self.assertTrue(math.isnan(original["a"]))

    def test_output_survives_strict_json(self):
        """`allow_nan=False` is what DRF's renderer and Postgres both enforce."""
        payload = json_safe({"ulr_rows": [PRODUCTION_ROW]})
        self.assertNotIn("NaN", json.dumps(payload))
        json.dumps(payload, allow_nan=False)  # must not raise


class InputMetaPersistenceTests(TestCase):
    """The field itself must refuse to hand Postgres something it will reject."""

    def setUp(self):
        self.org = Organization.objects.create(name="Acme", slug="acme-jsonsafe")
        self.user = User.objects.create_user(username="json-safe-alice", password="x")

    def _job(self, meta):
        return Module1Job.objects.create(
            user=self.user,
            organization=self.org,
            job_type=Module1Job.JobType.MODULE2_ALLOCATE,
            status=Module1Job.Status.RUNNING,
            work_dir="module1_jobs/x",
            input_meta=meta,
        )

    def test_saving_ulr_rows_with_nan_no_longer_fails(self):
        job = self._job({"ulr_rows": [PRODUCTION_ROW]})
        job.refresh_from_db()
        row = job.input_meta["ulr_rows"][0]
        self.assertIsNone(row["exp_ratio"])
        self.assertIsNone(row["ri_percent"])
        self.assertIsNone(row["combined_ratio"])
        # The real numbers must come back exactly as they went in.
        self.assertEqual(row["comm_ratio"], -0.009365161576319028)
        self.assertEqual(row["gwp"], 1768576.0)
        self.assertEqual(row["reserving_class"], "Fire")

    def test_update_fields_save_path_is_covered(self):
        """tasks.py fails on `job.save(update_fields=[...])`, not on create."""
        job = self._job({})
        job.input_meta = {"ulr_rows": [PRODUCTION_ROW], "override_report": {"x": float("inf")}}
        job.save(update_fields=["input_meta"])
        job.refresh_from_db()
        self.assertIsNone(job.input_meta["ulr_rows"][0]["exp_ratio"])
        self.assertIsNone(job.input_meta["override_report"]["x"])

    def test_sensitivity_and_movement_shaped_payloads_are_covered(self):
        """The same landmine sat under `meta["sensitivity"]` and
        `meta["movement_warnings"]`; the field covers every key, not just ULR."""
        job = self._job({})
        job.input_meta = {
            "sensitivity": {"base": {"lic": {"total": float("nan")}}},
            "movement_warnings": {"reconciliation": {"max_abs_residual": float("nan")}},
        }
        job.save(update_fields=["input_meta"])
        job.refresh_from_db()
        self.assertIsNone(job.input_meta["sensitivity"]["base"]["lic"]["total"])
        self.assertIsNone(
            job.input_meta["movement_warnings"]["reconciliation"]["max_abs_residual"]
        )


# ---------------------------------------------------------------------------
# End-to-end: the real task, against the real database
# ---------------------------------------------------------------------------

import shutil
import tempfile
from pathlib import Path

from django.test import override_settings

from accounts.models import Permission, Role
from tenants.models import Membership

TASK_MEDIA_ROOT = tempfile.mkdtemp(prefix="sigma17-test-media-jsonsafe-")


@override_settings(MEDIA_ROOT=TASK_MEDIA_ROOT, SECURE_SSL_REDIRECT=False)
class AllocateTaskWithBlankRatiosTests(TestCase):
    """The whole failure, reproduced through the actual Celery task body.

    The client's Combined_Summary had blank `Exp Ratio` / `RI %` columns, so the
    engine emitted NaN for those and for the `Combined Ratio` derived from them.
    Blanking those two columns in the reference fixture reproduces exactly that.
    """

    FIXTURE = (
        Path(__file__).resolve().parents[2]
        / "benchmarks" / "fixtures" / "m2_allocate_ref" / "Combined_Summary.xlsx"
    )

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(TASK_MEDIA_ROOT, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        if not self.FIXTURE.is_file():
            self.skipTest("reference fixture not available")
        self.org = Organization.objects.create(name="JS", slug="js-e2e")
        self.user = User.objects.create_user("js-e2e", "js@example.com", "pw")
        self.user.profile.active_organization = self.org
        self.user.profile.save(update_fields=["active_organization"])
        role = Role.objects.create(name="ActuaryJS")
        perm, _ = Permission.objects.get_or_create(
            key="module2.run",
            defaults={"name": "module2.run", "module": "Processing", "description": ""},
        )
        role.permissions.add(perm)
        m = Membership.objects.create(user=self.user, organization=self.org, status="active")
        m.roles.add(role)

    def _workbook_with_blank_ratios(self) -> bytes:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(self.FIXTURE.read_bytes()))
        ws = wb["UW Summary"]
        headers = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
        for name in ("Exp Ratio", "RI %"):
            col = headers.get(name)
            if col is None:
                self.skipTest(f"fixture has no '{name}' column")
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=col).value = None
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_allocate_succeeds_when_ratio_columns_are_blank(self):
        from processing.tasks import run_module2_allocate_task
        from processing.utils import init_module2_allocate_job_dirs

        job = Module1Job.objects.create(
            user=self.user, organization=self.org,
            job_type=Module1Job.JobType.MODULE2_ALLOCATE,
        )
        job.work_dir = f"module1_jobs/{job.id}"
        job.save(update_fields=["work_dir"])
        combined_dir = init_module2_allocate_job_dirs(job)
        (combined_dir / "Combined_Summary.xlsx").write_bytes(
            self._workbook_with_blank_ratios()
        )
        job.input_meta = {"files": {}}
        job.save(update_fields=["input_meta"])

        run_module2_allocate_task(str(job.id))
        job.refresh_from_db()

        # Before the fix this was FAILED with "Module 2 processing failed. Check
        # workbook formats and required sheets" — the engine had already written
        # its workbook; only the json write blew up.
        self.assertEqual(job.status, Module1Job.Status.SUCCESS, job.error_message)
        rows = job.input_meta.get("ulr_rows") or []
        self.assertTrue(rows, "allocate produced no ULR rows")
        self.assertTrue(
            any(r.get("exp_ratio") is None for r in rows),
            "expected blank Exp Ratio to round-trip as null",
        )
        # And the run is genuinely usable, not just persisted.
        self.assertTrue(job.output_zip)
