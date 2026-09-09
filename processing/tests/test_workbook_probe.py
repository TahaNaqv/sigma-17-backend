"""The sheet probe reads xl/workbook.xml directly, so it must behave exactly
like a real parse would — and must never raise, because chaining treats
"unreadable" as "unknown" and a raised exception would 500 a picker page."""

import io
import zipfile

from django.test import SimpleTestCase
from openpyxl import Workbook

from processing.services.workbook_probe import missing_sheets, sheet_names_from_xlsx_bytes


def _xlsx(sheets, hidden=()) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name in sheets:
        ws = wb.create_sheet(title=name)
        ws.append(["a"])
        if name in hidden:
            ws.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class SheetProbeTests(SimpleTestCase):
    def test_returns_sheet_names_in_workbook_order(self):
        raw = _xlsx(["IBNR Summary", "ULAE-RA", "Discount Rate"])
        self.assertEqual(
            sheet_names_from_xlsx_bytes(raw),
            ["IBNR Summary", "ULAE-RA", "Discount Rate"],
        )

    def test_matches_openpyxl_on_a_realistic_book(self):
        names = ["Combined Summary", "LIC (OS) Summary", "UPR Run-Off", "UW Summary"]
        raw = _xlsx(names)
        from openpyxl import load_workbook

        self.assertEqual(
            sheet_names_from_xlsx_bytes(raw),
            load_workbook(io.BytesIO(raw), read_only=True).sheetnames,
        )

    def test_hidden_sheets_are_reported(self):
        """pandas reads a hidden sheet by name, so the engine can consume one;
        omitting it here would under-report what a workbook satisfies."""
        raw = _xlsx(["Visible", "ULAE-RA"], hidden=["ULAE-RA"])
        self.assertIn("ULAE-RA", sheet_names_from_xlsx_bytes(raw))

    def test_names_needing_xml_escaping_round_trip(self):
        raw = _xlsx(["P&L (Gross)", "A<B"])
        self.assertEqual(sheet_names_from_xlsx_bytes(raw), ["P&L (Gross)", "A<B"])

    # -- every failure mode returns [], never raises -------------------------

    def test_not_a_zip_returns_empty(self):
        self.assertEqual(sheet_names_from_xlsx_bytes(b"\xd0\xcf\x11\xe0legacy .xls"), [])

    def test_empty_input_returns_empty(self):
        self.assertEqual(sheet_names_from_xlsx_bytes(b""), [])

    def test_zip_without_workbook_part_returns_empty(self):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("not/a/workbook.txt", "hello")
        self.assertEqual(sheet_names_from_xlsx_bytes(buf.getvalue()), [])

    def test_malformed_workbook_xml_returns_empty(self):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("xl/workbook.xml", "<workbook><sheets>truncated")
        self.assertEqual(sheet_names_from_xlsx_bytes(buf.getvalue()), [])

    def test_truncated_archive_returns_empty(self):
        self.assertEqual(sheet_names_from_xlsx_bytes(_xlsx(["A"])[:40]), [])


class MissingSheetsTests(SimpleTestCase):
    def test_reports_in_required_order_not_alphabetical(self):
        self.assertEqual(
            missing_sheets(["IBNR Summary"], ("IBNR Summary", "ULAE-RA", "Discount Rate")),
            ["ULAE-RA", "Discount Rate"],
        )

    def test_nothing_missing_when_superset(self):
        self.assertEqual(missing_sheets(["A", "B", "C"], ("A", "C")), [])

    def test_comparison_is_exact(self):
        """pandas resolves a sheet by exact name, so a case variant really is
        missing and must be reported rather than quietly accepted."""
        self.assertEqual(missing_sheets(["ulae-ra"], ("ULAE-RA",)), ["ULAE-RA"])
