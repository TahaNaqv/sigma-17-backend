"""Data checks through the reserving task and the API: recorded on every run, never able to
fail a run, explained per check by the company, exported as the client's tables."""

import io
import shutil
import tempfile
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from django.contrib.auth.models import User
from django.test import TestCase, override_settings
from openpyxl import load_workbook
from rest_framework.test import APIClient

from processing.models import DataCheckExplanation, Module1Job
from processing.services.data_checks_store import DATA_CHECKS_FILENAME
from processing.tests.test_data_checks import _clean_book
from processing.tests.test_preflight_api import _give_role
from tenants.models import Organization

TEST_MEDIA_ROOT = tempfile.mkdtemp(prefix="sigma17-test-media-dc-")
EOP = "31-12-2024"


def _xlsx(frame: pd.DataFrame) -> bytes:
    buf = io.BytesIO()
    frame.to_excel(buf, index=False)
    return buf.getvalue()


def _fake_engine(*args, **kwargs):
    """Stands in for run_generate_summary: the reserve itself is not under test here."""
    out_dir = Path(args[7])
    pd.DataFrame({"x": [1]}).to_excel(out_dir / "Combined_Summary.xlsx", index=False)


@override_settings(MEDIA_ROOT=TEST_MEDIA_ROOT, SECURE_SSL_REDIRECT=False)
class DataChecksTaskTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(TEST_MEDIA_ROOT, ignore_errors=True)

    def setUp(self):
        self.org = Organization.objects.create(name="DC", slug="dc")
        self.user = User.objects.create_user("dc", "dc@example.com", "pw")
        _give_role(self.user, "ActuaryDC", ["module1.run", "runhistory.view"], self.org)
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    def _run(self, book=None):
        from processing.tasks import run_module1_summary_task
        from processing.utils import init_job_work_dir, job_input_subdir

        premium, paid, os_frame = book or _clean_book()
        job = Module1Job.objects.create(
            user=self.user, organization=self.org, job_type=Module1Job.JobType.SUMMARY,
            input_meta={"exp_start": "01-01-2024", "exp_end": EOP, "bop": "01-01-2024", "eop": EOP},
        )
        job.work_dir = f"module1_jobs/{job.id}"
        job.save(update_fields=["work_dir"])
        init_job_work_dir(job)
        for kind, frame in (("premium", premium), ("claims_paid", paid), ("claims_os", os_frame)):
            (job_input_subdir(job, kind) / f"{kind}.xlsx").write_bytes(_xlsx(frame))
        with patch("processing.tasks.run_generate_summary", side_effect=_fake_engine):
            run_module1_summary_task(str(job.id))
        job.refresh_from_db()
        return job

    def _failing_book(self):
        premium, paid, os_frame = _clean_book()
        premium.loc[0, "POLICYENDDATE"] = premium.loc[0, "ISSUEDATE"] - pd.Timedelta(days=1)
        return premium, paid, os_frame

    # ── task ────────────────────────────────────────────────────────────────

    def test_every_run_records_the_report_and_ships_the_evidence_workbook(self):
        job = self._run(self._failing_book())
        self.assertEqual(job.status, Module1Job.Status.SUCCESS, job.error_message)
        report = job.input_meta["data_checks"]
        by_id = {c["id"]: c for c in report["checks"]}
        self.assertEqual(by_id["2.4"]["status"], "fail")
        self.assertEqual(by_id["2.4"]["sample"][0]["__source_file"], "premium.xlsx")
        self.assertEqual(report["valuation_date"], "2024-12-31")
        self.assertIn(DATA_CHECKS_FILENAME, job.output_artifacts)

    def test_a_failure_inside_the_checks_never_fails_the_reserve(self):
        with patch("processing.services.data_checks.run_data_checks", side_effect=RuntimeError("x")):
            job = self._run()
        self.assertEqual(job.status, Module1Job.Status.SUCCESS, job.error_message)
        self.assertIn("could not be run", job.input_meta["data_checks"]["error"])

    def test_format_is_compared_with_the_previous_run(self):
        first = self._run()
        premium, paid, os_frame = _clean_book()
        second = self._run((premium.drop(columns=["COMMISSIONAMOUNT"]), paid, os_frame))
        report = second.input_meta["data_checks"]
        self.assertEqual(report["compared_with"]["job_id"], str(first.id))
        check = next(c for c in report["checks"] if c["id"] == "1.2")
        self.assertEqual(check["status"], "fail")
        self.assertEqual(check["sample"][0]["column"], "COMMISSIONAMOUNT")

    # ── API ─────────────────────────────────────────────────────────────────

    def _url(self, job, suffix=""):
        return f"/api/module1/jobs/{job.id}/data-checks/{suffix}"

    def test_get_returns_report_explanations_and_evidence_file(self):
        job = self._run(self._failing_book())
        res = self.client.get(self._url(job))
        self.assertEqual(res.status_code, 200, res.content)
        body = res.json()
        self.assertEqual(body["evidence_file"], DATA_CHECKS_FILENAME)
        self.assertEqual(body["explanations"], {})
        self.assertEqual(len(body["report"]["checks"]), 21)

    def test_explanations_are_saved_updated_and_cleared(self):
        job = self._run(self._failing_book())
        url = self._url(job, "2.4/explanation/")
        res = self.client.put(url, {"text": "  Endorsements after expiry are valid.  "}, format="json")
        self.assertEqual(res.status_code, 200, res.content)
        saved = res.json()["explanations"]["2.4"]
        self.assertEqual(saved["text"], "Endorsements after expiry are valid.")
        self.assertEqual(saved["updated_by"], "dc")
        self.client.put(url, {"text": "Revised."}, format="json")
        self.assertEqual(DataCheckExplanation.objects.get(job=job, check_id="2.4").text, "Revised.")
        self.client.put(url, {"text": ""}, format="json")
        self.assertFalse(DataCheckExplanation.objects.filter(job=job).exists())

    def test_explanation_validation(self):
        job = self._run()
        self.assertEqual(
            self.client.put(self._url(job, "9.9/explanation/"), {"text": "x"}, format="json").status_code, 400
        )
        self.assertEqual(
            self.client.put(self._url(job, "2.4/explanation/"), {"text": "x" * 4001},
                            format="json").status_code, 400
        )
        self.assertEqual(
            self.client.put(self._url(job, "2.4/explanation/"), {}, format="json").status_code, 400
        )

    def test_writing_an_explanation_requires_module1_run(self):
        job = self._run()
        viewer = User.objects.create_user("viewer", "v@example.com", "pw")
        _give_role(viewer, "ViewerDC", ["runhistory.view"], self.org)
        client = APIClient()
        client.force_authenticate(viewer)
        res = client.put(self._url(job, "2.4/explanation/"), {"text": "x"}, format="json")
        self.assertEqual(res.status_code, 403)

    def test_previous_runs_explanations_are_offered_to_carry_forward(self):
        first = self._run(self._failing_book())
        DataCheckExplanation.objects.create(job=first, check_id="2.4", text="Valid.", updated_by=self.user)
        second = self._run(self._failing_book())
        body = self.client.get(self._url(second)).json()
        self.assertEqual(body["explanations"], {})
        self.assertEqual(body["previous_explanations"]["2.4"]["text"], "Valid.")
        self.assertEqual(body["previous_explanations"]["2.4"]["job_id"], str(first.id))

    def test_export_is_the_client_tables_with_the_explanations(self):
        job = self._run(self._failing_book())
        self.client.put(self._url(job, "2.4/explanation/"), {"text": "Valid."}, format="json")
        res = self.client.get(self._url(job, "export/"))
        self.assertEqual(res.status_code, 200)
        self.assertIn("Data_Checks_Report_31-12-2024.xlsx", res["Content-Disposition"])
        wb = load_workbook(io.BytesIO(res.content))
        ws = wb["Data Checks"]
        rows = {r[0].value: r for r in ws.iter_rows(min_row=5)}
        self.assertEqual(rows["2.4"][-1].value, "Valid.")

    def test_a_job_without_checks_is_a_clear_400(self):
        job = Module1Job.objects.create(
            user=self.user, organization=self.org, job_type=Module1Job.JobType.SUMMARY,
            status=Module1Job.Status.SUCCESS, input_meta={},
        )
        res = self.client.get(self._url(job))
        self.assertEqual(res.status_code, 400)
        self.assertIn("predates", res.json()["detail"])

    def test_preflight_previews_the_checks_before_a_run(self):
        premium, paid, os_frame = self._failing_book()
        files = {}
        for name, frame in (("premium", premium), ("claims_paid", paid), ("claims_os", os_frame)):
            upload = io.BytesIO(_xlsx(frame))
            upload.name = f"{name}.xlsx"
            files[name] = upload
        res = self.client.post("/api/module1/preflight/", {**files, "eop": EOP}, format="multipart")
        self.assertEqual(res.status_code, 200, res.content)
        checks = {c["id"]: c for c in res.json()["data_checks"]["checks"]}
        self.assertEqual(checks["2.4"]["status"], "fail")
        self.assertEqual(checks["2.6"]["status"], "pass")
        self.assertFalse(res.json()["would_block"], "data checks never gate")
        self.assertFalse(Module1Job.objects.filter(organization=self.org).exists())
