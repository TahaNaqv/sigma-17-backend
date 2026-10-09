"""The client's data checklist (Table A.6) — each check's rule, its record/amount basis, and
the not-run contract, on small synthetic books where the right answer is known by hand."""

import io

import pandas as pd
from django.test import SimpleTestCase
from openpyxl import load_workbook

from processing.services import data_checks as D
from processing.services.data_checks_workbook import (
    CLAIMS_SHEET,
    EXPLANATION_HEADER,
    PREMIUM_SHEET,
    SUMMARY_HEADER_ROW,
    SUMMARY_SHEET,
    build_data_checks_workbook,
    detail_sheet_name,
)

VALUATION = "31-12-2024"


def _premium_row(policy, treaty, amount, *, issue="2024-01-01", start="2024-01-01",
                 end="2024-12-31", cls="MOTOR", endorsement=1, **extra):
    row = {
        "POLICYNUMBER": policy, "ENDORSEMENTNUMBER": endorsement, "RESERVINGCLASS": cls,
        "POLICYCLASS": cls, "IFRSCLASS": cls, "RI_TREATY_TYPE": treaty, "PREMIUMAMOUNT": amount,
        "COMMISSIONAMOUNT": 0.0, "ISSUEDATE": issue, "POLICYSTARTDATE": start,
        "POLICYENDDATE": end,
    }
    row.update(extra)
    return row


def _premium(rows):
    frame = pd.DataFrame(rows)
    for col in ("ISSUEDATE", "POLICYSTARTDATE", "POLICYENDDATE"):
        frame[col] = pd.to_datetime(frame[col], errors="coerce")
    return frame


def _claim_row(claim, treaty, paid, *, loss="2024-03-01", reported="2024-03-05",
               payment="2024-04-01", cls="MOTOR"):
    return {
        "CLAIMNUMBER": claim, "RESERVINGCLASS": cls, "POLICYCLASS": cls, "HEADOFDAMAGE": "Payment",
        "RI_TREATY_TYPE": treaty, "AMOUNTPAID": paid, "AMOUNTRECOVERED": 0.0,
        "LOSSDATE": loss, "REPORTEDDATE": reported, "PAYMENTDATE": payment,
    }


def _claims(rows):
    frame = pd.DataFrame(rows)
    for col in ("LOSSDATE", "REPORTEDDATE", "PAYMENTDATE"):
        if col in frame:
            frame[col] = pd.to_datetime(frame[col], errors="coerce")
    return frame


def _os(rows):
    frame = _claims(rows).drop(columns=["PAYMENTDATE", "AMOUNTPAID"])
    frame["AMOUNTOUTSTANDING"] = 1000.0
    frame["As at"] = pd.Timestamp("2024-12-31")
    return frame


def _clean_book():
    premium = _premium([
        _premium_row("P1", "GROSS", 1000.0), _premium_row("P1", "RI", 300.0),
        _premium_row("P2", "GROSS", 500.0), _premium_row("P2", "RI", 100.0),
    ])
    paid = _claims([_claim_row("C1", "GROSS", 200.0), _claim_row("C1", "RI", 50.0)])
    os_frame = _os([_claim_row("C2", "GROSS", 0.0)])
    return premium, paid, os_frame


def _run(premium, paid, os_frame, **kw):
    kw.setdefault("valuation_date", VALUATION)
    return {r.id: r for r in D.run_data_checks(premium, paid, os_frame, **kw).results}


class CatalogueTests(SimpleTestCase):
    def test_reference_numbers_and_order_follow_the_client_checklist(self):
        self.assertEqual(
            D.check_ids(),
            ["1.1", "1.2", "1.3", "1.4", "2.1", "2.2", "2.3", "2.4", "2.5", "2.6", "2.7",
             "2.8", "2.9", "2.11", "2.12", "3.1", "3.2", "3.3", "3.4", "3.5", "3.6"],
        )

    def test_a_clean_book_passes_every_check_that_can_run(self):
        results = _run(*_clean_book())
        failed = {k: r.basis for k, r in results.items() if r.status == D.FAIL}
        self.assertEqual(failed, {})
        self.assertEqual(results["1.2"].status, D.NOT_RUN)  # no previous valuation


class PremiumCheckTests(SimpleTestCase):
    def test_nwp_of_a_failing_record_is_gross_minus_ceded(self):
        premium, paid, os_frame = _clean_book()
        premium = pd.concat([premium, _premium([
            _premium_row("P3", "GROSS", 800.0, issue="2025-02-01", end="2024-12-31"),
            _premium_row("P3", "RI", 200.0),
        ])], ignore_index=True)
        r = _run(premium, paid, os_frame)["2.4"]
        self.assertEqual(r.status, D.FAIL)
        self.assertEqual(r.records, 1)
        self.assertEqual(r.amount, 600.0)

    def test_date_checks_read_gross_rows_not_the_cession_rows(self):
        """RI rows carry the cession's dates; a late cession is not a policy defect."""
        premium, paid, os_frame = _clean_book()
        premium.loc[premium["RI_TREATY_TYPE"] == "RI", "ISSUEDATE"] = pd.Timestamp("2026-01-01")
        results = _run(premium, paid, os_frame)
        self.assertEqual(results["2.4"].status, D.PASS)
        self.assertEqual(results["2.6"].status, D.PASS)

    def test_expiry_on_the_issue_date_fails_but_a_claim_paid_on_the_loss_date_passes(self):
        premium, paid, os_frame = _clean_book()
        premium.loc[0, "POLICYENDDATE"] = premium.loc[0, "ISSUEDATE"]
        paid.loc[:, "PAYMENTDATE"] = paid["LOSSDATE"]
        results = _run(premium, paid, os_frame)
        self.assertEqual(results["2.4"].records, 1)
        self.assertEqual(results["3.4"].status, D.PASS)

    def test_a_missing_date_is_counted_once_not_again_in_the_comparisons(self):
        premium, paid, os_frame = _clean_book()
        premium.loc[0, "ISSUEDATE"] = pd.NaT
        results = _run(premium, paid, os_frame)
        self.assertEqual(results["2.1"].records, 1)
        self.assertEqual(results["2.4"].status, D.PASS)
        self.assertEqual(results["2.6"].status, D.PASS)

    def test_an_implausible_date_is_invalid(self):
        premium, paid, os_frame = _clean_book()
        premium["POLICYSTARTDATE"] = premium["POLICYSTARTDATE"].astype(object)
        premium.loc[0, "POLICYSTARTDATE"] = "not a date"
        premium.loc[2, "POLICYSTARTDATE"] = pd.Timestamp("1850-01-01")
        self.assertEqual(_run(premium, paid, os_frame)["2.2"].records, 2)

    def test_effective_after_valuation(self):
        premium, paid, os_frame = _clean_book()
        premium.loc[premium["POLICYNUMBER"] == "P2", "POLICYSTARTDATE"] = pd.Timestamp("2025-01-01")
        premium.loc[premium["POLICYNUMBER"] == "P2", "POLICYENDDATE"] = pd.Timestamp("2025-12-31")
        r = _run(premium, paid, os_frame)["2.7"]
        self.assertEqual((r.records, r.amount), (1, 400.0))

    def test_zero_gross_with_ceded_premium(self):
        premium, paid, os_frame = _clean_book()
        premium = pd.concat([premium, _premium([
            _premium_row("P4", "GROSS", 0.0), _premium_row("P4", "RI", 25.0),
        ])], ignore_index=True)
        r = _run(premium, paid, os_frame)["2.8"]
        self.assertEqual((r.records, r.amount), (1, -25.0))

    def test_net_larger_than_gross(self):
        premium, paid, os_frame = _clean_book()
        premium = pd.concat([premium, _premium([
            _premium_row("P5", "GROSS", 100.0), _premium_row("P5", "RI", -50.0),
        ])], ignore_index=True)
        r = _run(premium, paid, os_frame)["2.9"]
        self.assertEqual((r.records, r.amount), (1, 150.0))

    def test_blank_classes_and_endorsements(self):
        premium, paid, os_frame = _clean_book()
        premium.loc[0, "IFRSCLASS"] = "  "
        premium.loc[2, "ENDORSEMENTNUMBER"] = None
        results = _run(premium, paid, os_frame)
        self.assertEqual(results["2.11"].records, 1)
        self.assertEqual(results["2.12"].records, 1)

    def test_a_check_whose_column_is_absent_is_not_run_never_passed(self):
        premium, paid, os_frame = _clean_book()
        results = _run(premium.drop(columns=["ENDORSEMENTNUMBER"]), paid, os_frame)
        self.assertEqual(results["2.12"].status, D.NOT_RUN)
        self.assertIn("ENDORSEMENTNUMBER", results["2.12"].basis)

    def test_without_a_valuation_date_the_valuation_checks_do_not_run(self):
        results = _run(*_clean_book(), valuation_date=None)
        for cid in ("2.6", "2.7", "3.6"):
            self.assertEqual(results[cid].status, D.NOT_RUN, cid)


class ClaimsCheckTests(SimpleTestCase):
    def test_records_are_claims_and_amount_is_net_paid(self):
        premium, paid, os_frame = _clean_book()
        bad = _claims([_claim_row("C9", "GROSS", 900.0, reported="2024-02-01"),
                       _claim_row("C9", "RI", 400.0, reported="2024-02-01")])
        r = _run(premium, pd.concat([paid, bad], ignore_index=True), os_frame)["3.5"]
        self.assertEqual((r.records, r.rows, r.amount), (1, 2, 500.0))

    def test_outstanding_rows_are_checked_too(self):
        premium, paid, os_frame = _clean_book()
        os_frame.loc[0, "LOSSDATE"] = pd.Timestamp("2025-03-01")
        results = _run(premium, paid, os_frame)
        self.assertEqual(results["3.6"].records, 1)
        self.assertEqual(results["3.6"].amount, 0.0)  # nothing paid on that claim

    def test_payment_date_checks_ignore_the_outstanding_file(self):
        premium, paid, os_frame = _clean_book()
        results = _run(premium, paid, os_frame)
        self.assertEqual(results["3.2"].status, D.PASS)


class GeneralCheckTests(SimpleTestCase):
    def test_a_class_missing_from_an_input_is_reported(self):
        premium, paid, os_frame = _clean_book()
        premium = pd.concat([premium, _premium([_premium_row("P6", "GROSS", 10.0, cls="MARINE")])],
                            ignore_index=True)
        r = _run(premium, paid, os_frame)["1.1"]
        self.assertEqual(r.records, 1)
        self.assertEqual(r.sample[0]["class"], "MARINE")

    def test_format_is_compared_with_the_previous_valuation(self):
        premium, paid, os_frame = _clean_book()
        baseline = {"job_id": "prev", "eop": "31-12-2023",
                    "input_schema": D.run_data_checks(premium, paid, os_frame).input_schema}
        changed = premium.drop(columns=["COMMISSIONAMOUNT"]).assign(NEWCOL=1)
        changed["PREMIUMAMOUNT"] = changed["PREMIUMAMOUNT"].astype(str)
        r = _run(changed, paid, os_frame, baseline=baseline)["1.2"]
        changes = {(row["column"], row["change"]) for row in r.sample}
        self.assertIn(("COMMISSIONAMOUNT", "missing (was present previously)"), changes)
        self.assertIn(("NEWCOL", "new column"), changes)
        self.assertIn(("PREMIUMAMOUNT", "type changed: number -> text"), changes)
        unchanged = _run(premium, paid, os_frame, baseline=baseline)["1.2"]
        self.assertEqual(unchanged.status, D.PASS)

    def test_text_in_a_number_column(self):
        premium, paid, os_frame = _clean_book()
        premium["PREMIUMAMOUNT"] = premium["PREMIUMAMOUNT"].astype(object)
        premium.loc[0, "PREMIUMAMOUNT"] = "1,000"
        self.assertEqual(_run(premium, paid, os_frame)["1.3"].records, 1)

    def test_duplicates_count_extra_copies_and_ignore_provenance(self):
        premium, paid, os_frame = _clean_book()
        dup = pd.concat([premium, premium.iloc[[0, 0]]], ignore_index=True)
        dup[D.SOURCE_ROW] = range(len(dup))  # differs on every row; must not hide the copy
        self.assertEqual(_run(dup, paid, os_frame)["1.4"].records, 2)

    def test_a_check_that_raises_is_reported_and_the_rest_still_run(self):
        original = D.CATALOGUE

        def boom(cd, ctx):
            raise RuntimeError("bad input")

        D.CATALOGUE = (D.CheckDef("9.9", D.CATEGORY_GENERAL, "Broken", boom), *original)
        try:
            results = _run(*_clean_book())
        finally:
            D.CATALOGUE = original
        self.assertEqual(results["9.9"].status, D.NOT_RUN)
        self.assertIn("bad input", results["9.9"].basis)
        self.assertEqual(results["2.1"].status, D.PASS)

    def test_the_report_is_json_safe(self):
        import json

        premium, paid, os_frame = _clean_book()
        premium.loc[0, "POLICYENDDATE"] = premium.loc[0, "ISSUEDATE"]
        json.dumps(D.run_data_checks(premium, paid, os_frame, valuation_date=VALUATION).as_dict())


class WorkbookTests(SimpleTestCase):
    def _failing_report(self):
        premium, paid, os_frame = _clean_book()
        premium.loc[0, "POLICYENDDATE"] = premium.loc[0, "ISSUEDATE"]
        paid.loc[:, "REPORTEDDATE"] = pd.Timestamp("2024-01-01")
        premium[D.SOURCE_FILE], premium[D.SOURCE_ROW] = "premium.xlsx", range(2, 2 + len(premium))
        return D.run_data_checks(premium, paid, os_frame, valuation_date=VALUATION)

    def test_evidence_workbook_has_the_client_tables_and_a_sheet_per_failing_check(self):
        report = self._failing_report()
        wb = load_workbook(io.BytesIO(build_data_checks_workbook(report)))
        self.assertEqual(wb.sheetnames[:4], [SUMMARY_SHEET, PREMIUM_SHEET, CLAIMS_SHEET, "Basis"])
        self.assertIn(detail_sheet_name("2.4"), wb.sheetnames)
        self.assertIn(detail_sheet_name("3.5"), wb.sheetnames)
        detail = wb[detail_sheet_name("2.4")]
        header = [c.value for c in detail[2]]
        self.assertEqual(header[:2], [D.SOURCE_FILE, D.SOURCE_ROW])
        self.assertEqual(detail.cell(row=3, column=1).value, "premium.xlsx")
        premiums = [[c.value for c in r] for r in wb[PREMIUM_SHEET].iter_rows(min_row=5)]
        # Expiry equals issue = effective in this book, so both 2.4 and 2.5 fail.
        self.assertEqual(premiums, [["2.4", 1, 700], ["2.5", 1, 700]])
        # No explanations at run time: the evidence has no explanation column.
        self.assertNotIn(EXPLANATION_HEADER, [c.value for c in wb[SUMMARY_SHEET][SUMMARY_HEADER_ROW]])

    def test_report_workbook_from_the_stored_dict_carries_explanations(self):
        stored = self._failing_report().as_dict()
        body = build_data_checks_workbook(stored, explanations={"2.4": "Valid endorsements."})
        wb = load_workbook(io.BytesIO(body))
        self.assertNotIn(detail_sheet_name("2.4"), wb.sheetnames)
        ws = wb[SUMMARY_SHEET]
        header = [c.value for c in ws[SUMMARY_HEADER_ROW]]
        col = header.index(EXPLANATION_HEADER) + 1
        by_ref = {r[0].value: r[col - 1].value for r in ws.iter_rows(min_row=SUMMARY_HEADER_ROW + 1)}
        self.assertEqual(by_ref["2.4"], "Valid endorsements.")
